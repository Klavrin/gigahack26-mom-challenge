"""Approved minutes -> n8n delivery contract v2 (schemas/meeting-delivery-v2.schema.json):
approved state snapshot + a DOCX of the minutes. The transcript never leaves the backend."""
import base64
import datetime as dt
import io

from docx import Document
from docx.shared import Pt

from . import render

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _strings(items) -> list[str]:
    return [s.strip() for s in items if isinstance(s, str) and s.strip()]


def to_state(job: dict, mom: dict, ended_at: dt.datetime) -> dict:
    started_at = ended_at - dt.timedelta(seconds=float(job.get("audio_s") or 0))
    notes = [f"Rezumat: {mom['summary'].strip()}"] if (mom.get("summary") or "").strip() else []
    return {
        "meeting": {
            "id": job["id"],
            "title": ((mom.get("title") or "").strip() or "Proces-verbal")[:240],
            "type": job["meeting_type"],
            "started_at": started_at.isoformat(timespec="seconds"),
            "ended_at": ended_at.isoformat(timespec="seconds"),
        },
        "participants": _strings(mom.get("participants", [])),
        "topics": _strings(f"{t.get('topic', '').strip()}: {t.get('discussion', '').strip()}".strip(": ")
                           for t in mom.get("topics", [])),
        "decisions": _strings(d.get("decision", "") for d in mom.get("decisions", [])),
        "action_items": [
            {
                "description": a["task"].strip(),
                "owner": (a.get("owner") or "").strip() or None,
                "deadline": (a.get("deadline") or a.get("deadline_text") or "").strip() or None,
                "status": "open",
            }
            for a in mom.get("action_items", []) if (a.get("task") or "").strip()
        ],
        "open_questions": _strings(mom.get("open_questions", [])),
        "important_notes": notes,
    }


def minutes_docx(mom: dict, meeting_type: str, date: dt.date) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    label = render.TYPE_LABELS.get(meeting_type, (meeting_type, ""))[0]
    doc.add_heading(mom.get("title") or "Proces-verbal", level=0)
    doc.add_paragraph(f"{label} · {date.isoformat()}")

    doc.add_heading("Rezumat", level=1)
    doc.add_paragraph(mom.get("summary") or "")
    if mom.get("participants"):
        doc.add_paragraph("Participanți: " + ", ".join(mom["participants"]))

    doc.add_heading("Decizii", level=1)
    for d in mom.get("decisions", []) or [{"decision": "Nicio decizie înregistrată."}]:
        doc.add_paragraph(d.get("decision", ""), style="List Number")

    doc.add_heading("Sarcini", level=1)
    items = mom.get("action_items", [])
    if items:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ("Sarcină", "Responsabil", "Termen")):
            cell.text = text
        for a in items:
            row = table.add_row().cells
            row[0].text = a.get("task", "")
            row[1].text = a.get("owner") or "Nespecificat"
            row[2].text = a.get("deadline") or a.get("deadline_text") or "—"
    else:
        doc.add_paragraph("Nicio sarcină înregistrată.")

    if mom.get("topics"):
        doc.add_heading("Subiecte discutate", level=1)
        for t in mom["topics"]:
            p = doc.add_paragraph()
            p.add_run(f"{t.get('topic', '')}. ").bold = True
            p.add_run(t.get("discussion", ""))
    if mom.get("open_questions"):
        doc.add_heading("Întrebări deschise", level=1)
        for q in mom["open_questions"]:
            doc.add_paragraph(q, style="List Bullet")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Generat on-premise. Înregistrarea și transcrierea nu au părăsit rețeaua internă."
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def payload(job: dict, mom: dict, approved_by: str, attempt: int) -> dict:
    date = dt.date.fromisoformat(job["meeting_date"])
    now = dt.datetime.now().astimezone()
    ended_at = dt.datetime.fromtimestamp(job["created"]).astimezone()   # upload = meeting over
    docx = minutes_docx(mom, job["meeting_type"], date)
    return {
        "schema_version": "2",
        "delivery_id": f"{job['id']}-{attempt}",
        "state": to_state(job, mom, min(ended_at, now)),
        "approval": {
            "status": "approved",
            "approved_by": approved_by.strip() or "Reviewer (web app)",
            "approved_at": now.isoformat(timespec="seconds"),
        },
        "documents": [{
            "filename": f"{job['id']}-minutes.docx",
            "mime_type": DOCX_MIME,
            "data_base64": base64.b64encode(docx).decode("ascii"),
        }],
    }
